import os
import re
import time
from datetime import datetime, timedelta
from zoneinfo import ZoneInfo

import paramiko
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH

bogota_tz = ZoneInfo("America/Bogota")
now = datetime.now(bogota_tz)
user = os.environ["USER"]
password = os.environ["PASSWORD"]
date_str = f"{now.month}_{now.day}_{now.strftime('%y')}"

file_path = rf"C:\Users\Andres.Acosta\Downloads\iSeries Turnover {date_str}.docx"
print(file_path)


# 2. DEFINE ALL FUNCTIONS FIRST
def run_as400_cmd(command):
    system_name = "PROD1"
    user = user
    pwd = password

    try:
        ssh = paramiko.SSHClient()
        ssh.set_missing_host_key_policy(paramiko.AutoAddPolicy())
        ssh.connect(system_name, username=user, password=pwd, timeout=10)

        full_cmd = f'system -i "{command}"'
        stdin, stdout, stderr = ssh.exec_command(full_cmd)

        output = stdout.read().decode("utf-8").strip()
        error = stderr.read().decode("utf-8").strip()

        ssh.close()

        combined = output + "\n" + error
        return combined

    except Exception as e:
        return f"❌ Connection Failed: {e}"


def dataAddition(data, counter):
    if os.path.exists(file_path):
        doc = Document(file_path)
        print("OK it exists")

        # Get the current hour formatted as "HH:00" (e.g., "18:00")
        current_sharp_hour = datetime.now(bogota_tz).strftime("%H:00")
        print(f"Searching table for hour: {current_sharp_hour}")

        # ✅ FIXED: Select the specific table index correctly here
        target_table = doc.tables[3]
        target_table_2 = doc.tables[1]

        found = False
        for row_idx, row in enumerate(target_table.rows):
            for i, cell in enumerate(row.cells):
                if current_sharp_hour in cell.text:

                    # --- ACTION 1: Fill out target_table (Table 4) ---
                    if i + counter < len(row.cells):
                        row.cells[i + counter].text = data
                        for paragraph in row.cells[i + counter].paragraphs:
                            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                            for run in paragraph.runs:
                                run.bold = True
                        found = True
                        print(f"Found match! Wrote data to target_table Row {row_idx}.")

                    # --- ACTION 2: Fill out target_table_2 (Table 2) conditionally ---
                    if current_sharp_hour == "20:00":
                        # Ensure the second table actually has this row to avoid errors
                        if row_idx < len(target_table_2.rows):
                            row2 = target_table_2.rows[row_idx]

                            # Verify the offset index is safe for Table 2's structure
                            if i + counter < len(row2.cells):
                                row2.cells[i + counter].text = data
                                for paragraph in row2.cells[i + counter].paragraphs:
                                    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                                    for run in paragraph.runs:
                                        run.bold = True
                                print(
                                    f"Filled out matching row in target_table_2 at Row {row_idx}."
                                )
                        else:
                            print(
                                f"⚠️ Warning: target_table_2 does not have a Row {row_idx}."
                            )

        if not found:
            print(f"⚠️ Warning: Could not find '{current_sharp_hour}' in Table 4.")

        doc.save(file_path)
        print("Update successful! Changes will sync shortly.")
    else:
        print(f"Error: Could not find the file at {file_path}")


def wait_until_next_sharp_hour():
    """Calculates seconds until the next exact sharp hour and sleeps."""
    now = datetime.now(bogota_tz)
    # Calculate the next hour mark
    next_hour = (now + timedelta(hours=1)).replace(minute=0, second=0, microsecond=0)
    delay = (next_hour - now).total_seconds()

    print(
        f"💤💤💤💤 Sleeping for {delay/60:.1f} minutes until {next_hour.strftime('%H:%M')}..."
    )
    time.sleep(delay)


# --- CONTINUOUS BACKGROUND LOOP ---
if __name__ == "__main__":
    print("🚀 Background ASP Tracker Started.")

    while True:
        # 1. Wait until the clock hits XX:00 sharp
        wait_until_next_sharp_hour()

        # 2. Your original execution logic runs here on the dot
        print(
            f"\n--- Running Automated Check at {datetime.now(bogota_tz).strftime('%H:%M')} ---"
        )
        raw_report = run_as400_cmd("dspaspbrm")

        asp_pattern = r"\s+([*\w]+)\s+\d{5}.*?\s(\d+\.\d+)\s+\d+"
        matches = re.findall(asp_pattern, raw_report)

        if matches:
            counter = 1
            for name, usage in matches:
                print(f"Found usage: {usage}")
                dataAddition(usage, counter)
                counter += 1
        else:
            print("❌ Parsing failed.")
